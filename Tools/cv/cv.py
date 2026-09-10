from pathlib import Path
import uuid

from flask import Flask, request, send_file, send_from_directory
import os
import docx

app = Flask(__name__)
APP_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = APP_DIR / "generated"
OUTPUT_DIR.mkdir(exist_ok=True)

@app.route("/")
def home():
    return send_from_directory(APP_DIR, "index.html")


@app.route("/generate", methods=["POST"])
def generate():
    name = request.form.get("name")
    title = request.form.get("title")
    email = request.form.get("email")
    phone = request.form.get("phone")
    address = request.form.get("address")
    summary = request.form.get("summary", "")
    skills = request.form.get("skills", "")
    job_titles = request.form.getlist("job_title[]")
    companies = request.form.getlist("company[]")
    dates = request.form.getlist("dates[]")
    responsibilities = request.form.getlist("responsibilities[]")
    education = request.form.get("education", "")

    docx_path = OUTPUT_DIR / f"cv_{uuid.uuid4().hex}.docx"
    doc_docx = docx.Document()
    
    if name:
        doc_docx.add_heading(name, 0)
    if title:
        doc_docx.add_heading(title, level=2)
        
    if email or phone or address:
        contact_info = []
        if email: contact_info.append(email)
        if phone: contact_info.append(phone)
        if address: contact_info.append(address)
        doc_docx.add_paragraph(" | ".join(contact_info))
        
    if summary:
        doc_docx.add_heading('Professional Summary', level=2)
        doc_docx.add_paragraph(summary)
        
    if skills:
        doc_docx.add_heading('Skills', level=2)
        doc_docx.add_paragraph(skills)
        
    doc_docx.add_heading('EXPERIENCE', level=1)
    
    for i in range(len(job_titles)):
        if i < len(job_titles) and job_titles[i].strip():
            doc_docx.add_heading(job_titles[i], level=2)
            comp = companies[i] if i < len(companies) else ""
            date = dates[i] if i < len(dates) else ""
            doc_docx.add_paragraph(f"{comp} | {date}")
            
            if i < len(responsibilities):
                for item in responsibilities[i].split("\n"):
                    item = item.strip()
                    if item:
                        doc_docx.add_paragraph(item, style='List Bullet')
                        
    if education:
        doc_docx.add_heading('Education', level=2)
        doc_docx.add_paragraph(education)
        
    doc_docx.save(docx_path)

    return send_file(docx_path, as_attachment=True, download_name="cv.docx")


if __name__ == "__main__":
    app.run(debug=True)
